"""Bulk-create Sales Commission Rules from the client's commission master workbook.

Dry run (default, writes nothing):
	bench --site <site> execute blueline.scripts.import_commission_rules.run \
		--kwargs "{'path': '/path/to/file.xlsx'}"

Create the rules listed under WOULD CREATE (all-or-nothing, through normal validation):
	bench --site <site> execute blueline.scripts.import_commission_rules.run \
		--kwargs "{'path': '/path/to/file.xlsx', 'execute': True}"

Add 'include_catch_all': True to also create rules with neither customer nor item (they
match every otherwise-unmatched sale in that company, so they are held back by default).

Nothing is ever created except Sales Commission Rule: Sales Person, Customer and Item are
resolved by name with case and whitespace ignored (trim, collapse spaces, casefold), and
anything that fails to resolve — or resolves to more than one record — is reported, not
created. Several people may share one customer+item (a sale can pay several payees).
Service-based rows (not item sales) are held out entirely; they need their own design.
Rows whose item is on HELD_ITEMS are never created, regardless of what exists on the site.
"""

import re
from collections import defaultdict

import frappe
import openpyxl
from frappe.utils import flt

RULE_DOCTYPE = "Sales Commission Rule"
SHEETS = ("Sales Person", "Customer Commissions")
FIXED = "Fixed Amount Per Unit"
PERCENTAGE = "Percentage of Line Value"

# Fallback when the row's department isn't an existing Department record.
DEPARTMENT_SUFFIX_COMPANY = {
	"GIL": "General Innovations (Pvt) Ltd",
	"BL": "Blueline Enterprises Pvt Ltd",
}

COL_NAME = "sales person name"
COL_RATE = "commission rate"
COL_DEPT = "department"
COL_CUSTOMER = "id (customer / company name)"
COL_TARGET = "target amount rs. or % (targets)"
COL_ITEM = "item code (targets)"
COL_BANK_NAME = "bank account name"

# GI technical-team 40% structure: service-based, not item sales. Held out of this import.
SERVICE_ROWS = {"engineering commissioning", "visit and inspection", "repair service", "special approved sales commission"}

# Explicit, reviewed item mappings — not fuzzy matching. Applies only when BOTH the row
# reference and the workbook's exact cell text match; the target must be an existing Item.
MANUAL_ITEM_ALIASES = {
	("Sales Person!24", "KCE-53-12PAT1-SV"): "Savema KCE-53-12PAT1-SV 53mm",
}

# Deliberately deferred pending client confirmation. Matched against the workbook's item text
# (case/space-insensitive) BEFORE any item matching, so these rows are never created — even in
# execute mode, and even when a same-named Item exists on the site.
HELD_ITEMS = [
	"170xi4 - 300 DPI",
	"3811 DOD Water Based Black Ink",
	'BL 080 1/2"',
	"BL110",
	"BLZLWRFIPC 33mmX450M BLK",
	"GI WR3811BK - HP",
	"GI WR3913 BK - HP",
	"SAVEMA 53C Thermal Printer Head",
]

# Explicit, reviewed customer mappings — not fuzzy matching. Keyed by the workbook's exact
# customer text; applied only when the target exists on this site (otherwise normal matching
# runs and the alias is reported as not applied).
MANUAL_CUSTOMER_ALIASES = {
	"Pyramid Lanka (Pvt) Limited": "PYRAMID LANKA PVT LTD",
	"DHT Cement (Pvt) Ltd": "DHT Cement Pvt Ltd",
	"Ceylon Tobacco Company PLC": "Ceylon Tobacco Company",
	"Sunshine Consumer Lanka": "Sunshine Consumer Lanka Pvt Ltd",
}

RS_EACH = re.compile(r"^rs\.?\s*([\d,]+(?:\.\d+)?)\s*(?:each|per\s*unit)?$", re.I)
PERCENT = re.compile(r"^([\d.]+)\s*%$")


def _text(value):
	if value is None:
		return ""
	if isinstance(value, float) and value.is_integer():
		value = int(value)
	return str(value).replace("\xa0", " ").strip()


def _norm(value):
	return re.sub(r"\s+", " ", _text(value)).casefold()


class NameIndex:
	"""Name lookup over one or more fields, ignoring case and whitespace differences.
	Fields are tried in order (ID first); a normalized value shared by several records
	is reported as ambiguous rather than guessed.
	"""

	def __init__(self, doctype, fields):
		rows = frappe.get_all(doctype, fields=["name", *fields])
		self.fields = fields
		self.by_field = [defaultdict(set) for _ in fields]
		for row in rows:
			for i, field in enumerate(fields):
				if row.get(field):
					self.by_field[i][_norm(row[field])].add(row.name)

	def resolve(self, value):
		"""-> (name or None, note)"""
		for field, index in zip(self.fields, self.by_field):
			hits = index.get(_norm(value))
			if hits and len(hits) == 1:
				name = next(iter(hits))
				if field != "name":
					return name, f"matched on {field}: {value!r} -> {name!r}"
				if name != value:
					return name, f"case/space-normalized: {value!r} -> {name!r}"
				return name, ""
			if hits:
				return None, f"ambiguous: {len(hits)} records match ({', '.join(sorted(hits))})"
		return None, ""


def _read_rows(path):
	wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
	for sheet in SHEETS:
		if sheet not in wb.sheetnames:
			frappe.throw(f"Sheet {sheet!r} not found in {path}")
		rows = list(wb[sheet].iter_rows(values_only=True))
		header = {_norm(h): i for i, h in enumerate(rows[0]) if _text(h)}
		missing = {COL_NAME, COL_RATE, COL_DEPT, COL_CUSTOMER, COL_TARGET, COL_ITEM, COL_BANK_NAME} - set(header)
		if missing:
			frappe.throw(f"Sheet {sheet!r} is missing columns: {sorted(missing)}")
		prev = None
		for rownum, values in enumerate(rows[1:], start=2):
			get = lambda col: _text(values[header[col]]) if header[col] < len(values) else ""
			row = {
				"ref": f"{sheet}!{rownum}",
				"person": get(COL_NAME) or get(COL_BANK_NAME),
				"person_source": "Sales Person Name" if get(COL_NAME) else "Bank Account Name",
				"rate_cell": values[header[COL_RATE]],
				"target_cell": get(COL_TARGET),
				"department": get(COL_DEPT),
				"customer": get(COL_CUSTOMER),
				"item": get(COL_ITEM),
				"notes": [],
			}
			if not any(_text(v) for v in values):
				continue
			if not row["person"]:
				row["person"] = " / ".join(_text(v) for v in values if _text(v))
				row["person_source"] = "no name column filled; whole row shown"
			# Continuation row: same person as the row above with only a new item filled in.
			if (prev and row["person"] == prev["person"] and row["item"]
					and not (row["customer"] or row["department"] or row["target_cell"] or _text(row["rate_cell"]))):
				for key in ("rate_cell", "target_cell", "department", "customer"):
					row[key] = prev[key]
				row["notes"].append(f"customer/rate/department filled down from {prev['ref']}")
			prev = row
			yield row


def _parse_rate(row):
	"""-> (commission_type, rate, note) or (None, None, reason)"""
	rate_cell = row["rate_cell"]
	if _text(rate_cell):
		try:
			rate = float(_text(rate_cell).replace(",", "").rstrip("%"))
		except ValueError:
			return None, None, f"Commission Rate {rate_cell!r} is not a number"
		note = f"rate from Commission Rate column; Target column {row['target_cell']!r} ignored" if row["target_cell"] else ""
		return PERCENTAGE, rate, note
	target = row["target_cell"]
	if m := RS_EACH.match(target):
		return FIXED, flt(m.group(1).replace(",", "")), ""
	if m := PERCENT.match(target):
		return PERCENTAGE, flt(m.group(1)), ""
	if target:
		return None, None, f"no Commission Rate, and Target {target!r} is not a 'Rs. N Each' or 'N%' rate"
	return None, None, "no rate in Commission Rate or Target column"


def _resolve_company(department):
	if not department:
		return None, "no department"
	dept = frappe.db.get_value("Department", department, ["name", "company"], as_dict=True)
	if dept and dept.name == department and dept.company:
		return dept.company, ""
	suffix = department.rsplit("-", 1)[-1].strip() if "-" in department else ""
	if suffix in DEPARTMENT_SUFFIX_COMPANY:
		return DEPARTMENT_SUFFIX_COMPANY[suffix], f"department {department!r} not a Department record; company inferred from suffix {suffix!r}"
	return None, f"cannot infer company from department {department!r}"


def run(path, execute=False, include_catch_all=False):
	persons = NameIndex("Sales Person", ["name"])
	customers = NameIndex("Customer", ["name", "customer_name"])
	items = NameIndex("Item", ["name", "item_name"])

	unmatched = {"Sales Person": defaultdict(list), "Customer": defaultdict(list), "Item": defaultdict(list), "Company": defaultdict(list)}
	hints = {}
	no_rate, unresolved, candidates, services, aliases_applied = [], [], [], [], []
	held_items, customer_aliases_applied = [], []
	held_norm = {_norm(i) for i in HELD_ITEMS}

	for row in _read_rows(path):
		if row["item"] and _norm(row["item"]) in held_norm:
			held_items.append(row)
			continue
		if _norm(row["item"]) in SERVICE_ROWS:
			services.append(row)
			continue
		ctype, rate, rate_note = _parse_rate(row)
		if not ctype:
			no_rate.append((row, rate_note))
			continue
		if rate_note:
			row["notes"].append(rate_note)

		failed = []
		sales_person, note = persons.resolve(row["person"]) if row["person"] else (None, "blank")
		if not sales_person:
			failed.append("Sales Person")
			unmatched["Sales Person"][row["person"] or "<blank>"].append(row["ref"])
			hints[("Sales Person", row["person"])] = note
		elif note:
			row["notes"].append(f"sales person {note}")
		customer = None
		if row["customer"]:
			alias = MANUAL_CUSTOMER_ALIASES.get(row["customer"])
			target = frappe.db.get_value("Customer", alias, "name") if alias else None
			if alias and target == alias:
				customer = target
				note = f"MANUALLY ALIASED (explicit table, not a fuzzy match): {row['customer']!r} -> {target!r}"
				customer_aliases_applied.append((row["ref"], row["customer"], target))
			else:
				customer, note = customers.resolve(row["customer"])
			if not customer:
				failed.append("Customer")
				unmatched["Customer"][row["customer"]].append(row["ref"])
				hints[("Customer", row["customer"])] = note
			elif note:
				row["notes"].append(f"customer {note}")
		item_code = None
		if row["item"]:
			alias = MANUAL_ITEM_ALIASES.get((row["ref"], row["item"]))
			if alias:
				item_code = frappe.db.get_value("Item", alias, "name")
				if item_code == alias:
					note = f"MANUALLY ALIASED (explicit table, not a fuzzy match): {row['item']!r} -> {item_code!r}"
					aliases_applied.append((row['ref'], row['item'], item_code))
				else:
					item_code, note = None, f"manual alias target {alias!r} is not an existing Item"
			else:
				item_code, note = items.resolve(row["item"])
			if not item_code:
				failed.append("Item")
				unmatched["Item"][row["item"]].append(row["ref"])
				hints[("Item", row["item"])] = note
			elif note:
				row["notes"].append(f"item {note}")
		company, note = _resolve_company(row["department"])
		if not company:
			failed.append("Company")
			unmatched["Company"][row["department"] or "<blank>"].append(row["ref"])
			hints[("Company", row["department"])] = note
		elif note:
			row["notes"].append(note)

		if failed:
			unresolved.append((row, failed))
			continue
		candidates.append({
			"row": row, "sales_person": sales_person, "customer": customer or "", "item_code": item_code or "",
			"company": company, "commission_type": ctype, "rate": rate,
		})

	existing = frappe.get_all(RULE_DOCTYPE, fields=["name", "sales_person", "customer", "item_code", "company", "enabled", "commission_type", "rate"])
	existing_by_combo = defaultdict(list)
	for r in existing:
		existing_by_combo[(r.company, r.customer or "", r.item_code or "")].append(r)

	# One rule per sales_person + customer + item + company (mirrors Sales Commission Rule
	# validation). Different people on the same customer+item are separate payees, not a conflict.
	by_key = defaultdict(list)
	for c in candidates:
		by_key[(c["company"], c["customer"], c["item_code"], c["sales_person"])].append(c)
	payees_by_combo = defaultdict(set)
	for company, customer, item_code, sales_person in by_key:
		payees_by_combo[(company, customer, item_code)].add(sales_person)

	create, already, blocked, held = [], [], [], []
	for (company, customer, item_code, sales_person), group in by_key.items():
		combo = (company, customer, item_code)
		terms = {(c["commission_type"], c["rate"]) for c in group}
		if len(terms) > 1:
			blocked.append((group, f"file gives {sales_person} different rates for the same customer+item: {sorted(terms)}"))
			continue
		c = group[0]
		if len(group) > 1:
			c["row"]["notes"].append(f"duplicate rows collapsed: {', '.join(g['row']['ref'] for g in group[1:])}")
		same = [r for r in existing_by_combo[combo] if r.sales_person == sales_person]
		enabled_same = [r for r in same if r.enabled]
		if enabled_same:
			r = enabled_same[0]
			diff = "" if (r.commission_type, flt(r.rate)) == (c["commission_type"], c["rate"]) else \
				f" — NOTE existing is {r.commission_type} {flt(r.rate):g}, file says {c['commission_type']} {c['rate']:g}"
			already.append((c, f"{r.name}{diff}"))
			continue
		if same:
			c["row"]["notes"].append(f"a DISABLED identical rule exists ({', '.join(r.name for r in same)}); a new enabled one would be created")
		others = (payees_by_combo[combo] | {r.sales_person for r in existing_by_combo[combo] if r.enabled}) - {sales_person}
		if others:
			c["row"]["notes"].append(f"shared customer+item: also pays {', '.join(sorted(others))}")
		if not customer and not item_code:
			held.append(c)
			continue
		create.append(c)

	_report(path, execute, include_catch_all, create, already, blocked, held, services, aliases_applied,
	        held_items, customer_aliases_applied, unresolved, no_rate, unmatched, hints)

	if not execute:
		print("\nDRY RUN — nothing was written.")
		frappe.db.rollback()
		return

	to_insert = create + (held if include_catch_all else [])
	created = []
	try:
		for c in to_insert:
			doc = frappe.get_doc({
				"doctype": RULE_DOCTYPE, "sales_person": c["sales_person"], "customer": c["customer"] or None,
				"item_code": c["item_code"] or None, "company": c["company"], "commission_type": c["commission_type"],
				"rate": c["rate"], "enabled": 1,
			}).insert()
			created.append(doc.name)
	except Exception:
		frappe.db.rollback()
		print(f"\nEXECUTE FAILED at {c['row']['ref']} — rolled back, nothing was created.")
		raise
	frappe.db.commit()
	print(f"\nEXECUTED — created {len(created)} Sales Commission Rule(s): {', '.join(created)}")


def _line(c):
	cust = c["customer"] or "<any customer>"
	item = c["item_code"] or "<any item>"
	rate = f"{c['rate']:g}%" if c["commission_type"] == PERCENTAGE else f"Rs. {c['rate']:g}/unit"
	notes = f"   [{'; '.join(c['row']['notes'])}]" if c["row"]["notes"] else ""
	return f"  {c['row']['ref']:<26} {c['sales_person']} | {cust} | {item} | {c['company']} | {rate}{notes}"


def _report(path, execute, include_catch_all, create, already, blocked, held, services, aliases_applied,
            held_items, customer_aliases_applied, unresolved, no_rate, unmatched, hints):
	p = print
	p(f"SALES COMMISSION RULE IMPORT — {'EXECUTE' if execute else 'DRY RUN'}")
	p(f"File: {path}")
	p(f"Totals: would create {len(create)} | already exist {len(already)} | blocked (same person, conflicting rates) {sum(len(g) for g, _ in blocked)} rows"
	  f" | held pending confirmation {len(held_items)} | catch-all held {len(held)} | service-based held {len(services)}"
	  f" | unresolved {len(unresolved)} | no usable rate {len(no_rate)}")

	p(f"\n=== MANUAL ITEM ALIASES ({len(MANUAL_ITEM_ALIASES)} defined, {len(aliases_applied)} applied) — explicit mappings, not fuzzy matches")
	for (ref, text), target in MANUAL_ITEM_ALIASES.items():
		applied = any(a[0] == ref and a[1] == text for a in aliases_applied)
		p(f"  {ref:<26} {text!r} -> {target!r}  [{'APPLIED' if applied else 'NOT APPLIED (row/text not found or target missing)'}]")

	p(f"\n=== MANUAL CUSTOMER ALIASES ({len(MANUAL_CUSTOMER_ALIASES)} defined, {len(customer_aliases_applied)} row(s) aliased) — explicit mappings, not fuzzy matches")
	for text, target in MANUAL_CUSTOMER_ALIASES.items():
		refs = [a[0] for a in customer_aliases_applied if a[1] == text]
		if refs:
			status = f"APPLIED to {', '.join(refs)}"
		elif frappe.db.get_value("Customer", target, "name") != target:
			status = "NOT APPLIED (target is not a Customer on this site; normal matching used)"
		else:
			status = "NOT APPLIED (no row with this exact customer text)"
		p(f"  {text!r} -> {target!r}  [{status}]")

	p(f"\n=== HELD — PENDING CONFIRMATION (explicit hold list) ({len(held_items)}) — NEVER created, even in execute mode")
	for row in held_items:
		p(f"  {row['ref']:<26} item={row['item']!r} | person={row['person']!r} | customer={row['customer']!r}")

	p(f"\n=== WOULD CREATE ({len(create)}) — row | sales_person | customer | item_code | company | rate")
	for c in create:
		p(_line(c))

	p(f"\n=== SKIPPED — IDENTICAL ENABLED RULE ALREADY EXISTS ({len(already)})")
	for c, why in already:
		p(_line(c) + f"  -> {why}")

	p(f"\n=== BLOCKED — SAME PERSON, CONFLICTING RATES ({len(blocked)})")
	for group, why in blocked:
		p(f"  * {why}")
		for c in group:
			p("  " + _line(c))

	p(f"\n=== HELD — CATCH-ALL RULES (no customer, no item) ({len(held)}) — "
	  + ("WILL be created" if include_catch_all else "NOT created unless include_catch_all=True"))
	p("    These match EVERY sale in the company that no narrower rule matches.")
	for c in held:
		p(_line(c))

	p(f"\n=== HELD — SERVICE-BASED, NOT ITEM-BASED ({len(services)}) — never imported; needs its own commission design")
	for row in services:
		p(f"  {row['ref']:<26} person={row['person']!r} service={row['item']!r} dept={row['department']!r} "
		  f"commission rate={_text(row['rate_cell'])!r}%")

	p(f"\n=== NOT BUILT — UNRESOLVED ({len(unresolved)} rows)")
	for row, failed in unresolved:
		p(f"  {row['ref']:<26} failed: {', '.join(failed)} | person={row['person']!r} ({row['person_source']}) "
		  f"customer={row['customer']!r} item={row['item']!r} dept={row['department']!r}")

	p(f"\n=== SKIPPED — NO USABLE RATE ({len(no_rate)} rows)")
	for row, why in no_rate:
		p(f"  {row['ref']:<26} person={row['person']!r} customer={row['customer']!r} item={row['item']!r} — {why}")

	for kind, title in (("Sales Person", "UNMATCHED SALES PERSON NAMES"), ("Item", "UNMATCHED ITEMS"),
	                    ("Customer", "UNMATCHED CUSTOMERS"), ("Company", "UNRESOLVED COMPANY / DEPARTMENT")):
		p(f"\n=== {title} ({len(unmatched[kind])} distinct) — nothing is created for these")
		for value, refs in sorted(unmatched[kind].items()):
			hint = hints.get((kind, value if value != "<blank>" else ""), "")
			hint = f"   ({hint})" if hint else ""
			p(f"  {value!r}  — rows {', '.join(refs)}{hint}")
